"""World Bank monthly benchmark capture and atomic Parquet publication."""

import hashlib
import json
import os
import re
import tempfile
import time
import urllib.request
import xml.etree.ElementTree as ET
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from coffee_cocoa_platform.paths import ProjectPaths
from coffee_cocoa_platform.runtime import coordinated_capture
from coffee_cocoa_platform.transport import retryable

SOURCE_URL = (
    "https://thedocs.worldbank.org/en/doc/"
    "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/"
    "CMO-Historical-Data-Monthly.xlsx"
)
DATASET = "world_bank_pink_sheet_monthly"
SCHEMA_VERSION = "1"
MAX_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
SERIES = {
    "Cocoa": "cocoa",
    "Coffee, Arabica": "coffee_arabica",
    "Coffee, Robusta": "coffee_robusta",
}
DESCRIPTION_PREFIXES = {
    "Cocoa (ICCO)": "cocoa",
    "Coffee, Arabica (ICO)": "coffee_arabica",
    "Coffee, Robusta (ICO)": "coffee_robusta",
}
XML_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
XML_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PERIOD_RE = re.compile(r"^(\d{4})M(0[1-9]|1[0-2])$")
DECIMAL_SCALE = Decimal("0.00000001")
MISSING_MARKERS = {"..", "n.a.", "N/A"}

PRICE_SCHEMA = pa.schema(
    [
        pa.field("source_dataset", pa.string(), nullable=False),
        pa.field("series_id", pa.string(), nullable=False),
        pa.field("period_month", pa.date32(), nullable=False),
        pa.field("source_value_text", pa.string()),
        pa.field("price_usd_per_kg", pa.decimal128(20, 8)),
        pa.field("source_capture_id", pa.string(), nullable=False),
        pa.field("source_status", pa.string()),
    ]
)

FORECAST_CAPTURE_SCHEMA = pa.schema(
    [
        pa.field("source_capture_id", pa.string(), nullable=False),
        pa.field("source_update_text", pa.string()),
        pa.field("source_update_date", pa.date32()),
        pa.field("retrieved_at_utc", pa.timestamp("us", tz="UTC"), nullable=False),
        pa.field(
            "first_seen_capture_at_utc", pa.timestamp("us", tz="UTC"), nullable=False
        ),
    ]
)


def write_forecast_capture_metadata(directory: Path, captures: dict) -> None:
    """Publish capture provenance separately from immutable price observations."""
    output = directory / "forecast_capture_metadata.parquet"
    previously_seen = {}
    if output.exists():
        previous = pq.read_table(output, schema=FORECAST_CAPTURE_SCHEMA)
        previously_seen = {
            row["source_capture_id"]: row["first_seen_capture_at_utc"]
            for row in previous.to_pylist()
        }
    rows = []
    for capture_id, record in sorted(captures.items()):
        retrieved = datetime.fromisoformat(record["retrieved_at_utc"])
        first_seen = datetime.fromisoformat(record["first_seen_capture_at_utc"])
        if retrieved.tzinfo is None or first_seen.tzinfo is None:
            raise PriceSourceError("Capture timestamps must have timezones")
        if capture_id in previously_seen:
            first_seen = min(first_seen, previously_seen[capture_id])
        update_text = record.get("source_update")
        update_date = _update_date(update_text) if update_text else None
        rows.append(
            {
                "source_capture_id": capture_id,
                "source_update_text": update_text,
                "source_update_date": date.fromisoformat(update_date)
                if update_date
                else None,
                "retrieved_at_utc": retrieved.astimezone(UTC),
                "first_seen_capture_at_utc": first_seen.astimezone(UTC),
            }
        )
    directory.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=directory, suffix=".parquet", delete=False
    ) as tmp:
        staged = Path(tmp.name)
    try:
        pq.write_table(
            pa.Table.from_pylist(rows, schema=FORECAST_CAPTURE_SCHEMA),
            staged,
            compression="zstd",
        )
        os.replace(staged, output)
    finally:
        staged.unlink(missing_ok=True)


class PriceSourceError(ValueError):
    """The capture cannot be trusted for publication."""


def _month(value: str) -> date:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", value):
        raise PriceSourceError(f"Invalid month bound: {value}")
    return date.fromisoformat(value + "-01")


def _update_date(label: str) -> str | None:
    if not label.startswith("Updated on "):
        return None
    try:
        # Publisher labels name a date, not an instant or timezone.
        return (
            datetime.strptime(label.removeprefix("Updated on "), "%B %d, %Y")
            .replace(tzinfo=UTC)
            .date()
            .isoformat()
        )
    except ValueError:
        return None


def _sheet_xml(archive: ZipFile, sheet_name: str) -> bytes:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.attrib["Id"]: r.attrib["Target"] for r in relationships}
    for sheet in workbook.findall(f".//{{{XML_MAIN}}}sheet"):
        if sheet.attrib["name"] == sheet_name:
            target = targets[sheet.attrib[f"{{{XML_REL}}}id"]]
            member = target.lstrip("/") if target.startswith("/") else "xl/" + target
            return archive.read(member)
    raise PriceSourceError(f"Missing sheet: {sheet_name}")


def _raw_cells(data: bytes) -> dict[str, tuple[str | None, str | None]]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 100
                or sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED_BYTES
            ):
                raise PriceSourceError("Workbook exceeds decompressed size limit")
            xml = _sheet_xml(archive, "Monthly Prices")
        root = ET.fromstring(xml)
    except (BadZipFile, ET.ParseError, KeyError) as exc:
        raise PriceSourceError("Invalid XLSX structure") from exc
    cells = {}
    for cell in root.findall(
        f".//{{{XML_MAIN}}}sheetData/{{{XML_MAIN}}}row/{{{XML_MAIN}}}c"
    ):
        value = cell.find(f"{{{XML_MAIN}}}v")
        formula = cell.find(f"{{{XML_MAIN}}}f")
        cells[cell.attrib["r"]] = (
            value.text if value is not None else None,
            "formula" if formula is not None else cell.attrib.get("t"),
        )
    return cells


def _decimal(raw: str | None, kind: str | None, address: str) -> Decimal | None:
    if kind == "formula":
        raise PriceSourceError(f"Formula in {address}")
    if raw is None or raw == "":
        return None
    if kind not in (None, "n"):
        raise PriceSourceError(f"Non-numeric value in {address}")
    try:
        value = Decimal(raw)
        if not value.is_finite():
            raise InvalidOperation
        rounded = value.quantize(DECIMAL_SCALE, rounding=ROUND_HALF_EVEN)
        if len(rounded.as_tuple().digits) - rounded.as_tuple().exponent > 28:
            raise InvalidOperation
        if abs(rounded) >= Decimal(1000000000000):
            raise InvalidOperation
        return rounded
    except InvalidOperation as exc:
        raise PriceSourceError(f"Invalid decimal in {address}") from exc


def parse_workbook(
    data: bytes, start: str, end: str, capture_id: str
) -> tuple[pa.Table, dict]:
    """Validate the named worksheet and select inclusive reference months."""
    if len(data) > MAX_BYTES:
        raise PriceSourceError("Workbook exceeds download size limit")
    first, last = _month(start), _month(end)
    if first > last or (last.year - first.year) * 12 + last.month - first.month > 179:
        raise PriceSourceError("Month range must be ordered and at most 180 months")
    raw_cells = _raw_cells(data)
    try:
        book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False)
        sheet = book["Monthly Prices"]
        description_sheet = book["Description"]
        descriptions = {}
        for description_row in description_sheet.iter_rows(values_only=True):
            for value in description_row:
                if not isinstance(value, str):
                    continue
                for prefix, series_id in DESCRIPTION_PREFIXES.items():
                    if value.startswith(prefix):
                        if series_id in descriptions:
                            raise PriceSourceError(
                                f"Duplicate description for {series_id}"
                            )
                        descriptions[series_id] = value
        if set(descriptions) != set(SERIES.values()):
            raise PriceSourceError("Missing selected series description")
        rows = sheet.iter_rows()
        previous = None
        for header_row, candidate in enumerate(rows, start=1):
            if header_row > 20:
                raise PriceSourceError("Missing selected series headers")
            if all(
                candidate[col - 1].value == label
                for col, label in zip((12, 13, 14), SERIES)
            ):
                break
            previous = candidate
        else:
            raise PriceSourceError("Missing selected series headers")
        if previous is None:
            raise PriceSourceError("Missing update label row")
        units = next(rows, None)
        if units is None or any(
            units[col - 1].value != "($/kg)" for col in (12, 13, 14)
        ):
            raise PriceSourceError("Selected series units changed")
        update_text = str(previous[0].value or "")
        if not update_text.strip():
            raise PriceSourceError("Missing source update label")
        records = []
        seen = set()
        observed = {series_id: [] for series_id in SERIES.values()}
        represented = set()
        for row_number, row in enumerate(rows, start=header_row + 2):
            period_value = row[0].value
            selected = row[11:14]
            if period_value is None and all(cell.value is None for cell in selected):
                continue
            if not isinstance(period_value, str) or not (
                match := PERIOD_RE.fullmatch(period_value)
            ):
                raise PriceSourceError(f"Invalid period in A{row_number}")
            period = date(int(match.group(1)), int(match.group(2)), 1)
            if period in seen:
                raise PriceSourceError(f"Duplicate period: {period_value}")
            seen.add(period)
            if not first <= period <= last:
                continue
            represented.add(period)
            for col, series_id, cell in zip(("L", "M", "N"), SERIES.values(), selected):
                address = f"{col}{row_number}"
                raw, kind = raw_cells.get(address, (None, None))
                if (
                    kind != "formula"
                    and isinstance(cell.value, str)
                    and cell.value.strip() in MISSING_MARKERS
                ):
                    raw, price = cell.value, None
                else:
                    if cell.value is not None and raw is None and kind != "formula":
                        raise PriceSourceError(
                            f"Non-numeric or missing XML value in {address}"
                        )
                    price = _decimal(raw, kind, address)
                if price is not None:
                    observed[series_id].append(period)
                records.append(
                    (DATASET, series_id, period, raw, price, capture_id, None)
                )
        book.close()
    except (BadZipFile, KeyError, OSError) as exc:
        raise PriceSourceError("Invalid XLSX workbook") from exc
    except ValueError as exc:
        if isinstance(exc, PriceSourceError):
            raise
        raise PriceSourceError("Invalid Excel cell") from exc
    if not records:
        raise PriceSourceError("No selected periods in workbook")
    columns = list(zip(*records))
    table = pa.Table.from_arrays(
        [
            pa.array(values, type=field.type)
            for values, field in zip(columns, PRICE_SCHEMA)
        ],
        schema=PRICE_SCHEMA,
    )
    requested = (last.year - first.year) * 12 + last.month - first.month + 1
    coverage = {
        "series_descriptions": descriptions,
        "source_update_date": _update_date(update_text),
        "requested_month_count": requested,
        "represented_month_count": len(represented),
        "missing_month_count": requested - len(represented),
        "null_price_count": table.column("price_usd_per_kg").null_count,
        "latest_observed_by_series": {
            series: max(months).isoformat() if months else None
            for series, months in observed.items()
        },
    }
    return table, {"source_update_text": update_text, **coverage}


def _fetch_workbook(url: str, budget: dict) -> tuple[bytes, str, datetime, str | None]:
    """Consume a shared byte/time allowance across transport attempts."""
    retrieved = datetime.now(UTC)
    request = urllib.request.Request(
        url, headers={"User-Agent": "coffee-cocoa-platform/0.0.0"}
    )
    timeout = min(15, max(0.1, budget["deadline"] - time.monotonic()))
    with urllib.request.urlopen(request, timeout=timeout) as response:
        if response.status != 200:
            raise PriceSourceError(f"HTTP {response.status}")
        effective_url = response.geturl()
        content_type = response.headers.get("Content-Type")
        chunks = []
        while True:
            try:
                chunk = response.read(min(65536, budget["remaining_bytes"] + 1))
            except Exception as exc:
                partial = getattr(exc, "partial", b"")
                if isinstance(partial, bytes):
                    budget["remaining_bytes"] -= len(partial)
                raise
            budget["remaining_bytes"] -= len(chunk)
            if budget["remaining_bytes"] < 0 or time.monotonic() > budget["deadline"]:
                raise PriceSourceError(
                    "Live download exceeded cumulative byte or time limit"
                )
            if not chunk:
                break
            chunks.append(chunk)
    return b"".join(chunks), effective_url, retrieved, content_type


def fetch_workbook(url: str = SOURCE_URL):
    """At most three attempts within one 2 MiB/60-second allowance."""
    budget = {"remaining_bytes": MAX_BYTES, "deadline": time.monotonic() + 60}
    for attempt in range(3):
        if budget["remaining_bytes"] <= 0 or time.monotonic() >= budget["deadline"]:
            raise PriceSourceError(
                "Live download exhausted cumulative byte or time limit"
            )
        try:
            return _fetch_workbook(url, budget)
        except Exception as exc:
            if not retryable(exc) or attempt == 2:
                raise
            time.sleep(0.25 * (attempt + 1))


@coordinated_capture
def publish_workbook(
    data: bytes,
    *,
    start: str,
    end: str,
    paths: ProjectPaths,
    source_url: str,
    effective_url: str,
    retrieved_at: datetime,
    fixture: bool,
    content_type: str | None = None,
) -> dict:
    """Validate fully, then replace the current Parquet file atomically."""
    checksum = hashlib.sha256(data).hexdigest()
    table, coverage = parse_workbook(data, start, end, checksum)
    paths.raw.mkdir(parents=True, exist_ok=True)
    paths.parquet.mkdir(parents=True, exist_ok=True)
    raw_path = paths.raw / f"world_bank_monthly_{checksum}.xlsx"
    if raw_path.exists():
        if hashlib.sha256(raw_path.read_bytes()).hexdigest() != checksum:
            raise PriceSourceError("Existing raw capture checksum mismatch")
    else:
        with tempfile.NamedTemporaryFile(
            dir=paths.raw, suffix=".xlsx", delete=False
        ) as tmp:
            tmp.write(data)
            staged_raw = Path(tmp.name)
        try:
            os.replace(staged_raw, raw_path)
        finally:
            staged_raw.unlink(missing_ok=True)
    output = paths.parquet / "benchmark_prices.parquet"
    manifest_path = paths.parquet / "benchmark_prices.json"
    manifest = {
        "source_dataset": DATASET,
        "publisher": "World Bank",
        "terms_url": "https://www.worldbank.org/ext/en/legal/terms-conditions/datasets",
        "source_url": source_url,
        "effective_url": effective_url,
        "content_type": content_type,
        "retrieved_at_utc": retrieved_at.astimezone(UTC).isoformat(),
        "source_capture_id": checksum,
        "sha256": checksum,
        "input_start": start,
        "input_end": end,
        "schema_version": SCHEMA_VERSION,
        "row_count": table.num_rows,
        "byte_count": len(data),
        "http_status": None if fixture else 200,
        "observed_row_count": table.num_rows - coverage["null_price_count"],
        "status_capability": "not_provided",
        "fixture": fixture,
        "raw_path": str(raw_path),
        "parquet_path": str(output),
        **coverage,
    }
    with tempfile.NamedTemporaryFile(
        dir=paths.parquet, suffix=".parquet", delete=False
    ) as tmp:
        staged = Path(tmp.name)
    try:
        pq.write_table(table, staged, compression="zstd")
        if pq.read_metadata(staged).num_rows != table.num_rows:
            raise PriceSourceError("Staged Parquet row count mismatch")
        os.replace(staged, output)
    finally:
        staged.unlink(missing_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=paths.parquet, mode="w", encoding="utf-8", delete=False
    ) as tmp:
        json.dump(manifest, tmp, indent=2, sort_keys=True)
        staged_manifest = Path(tmp.name)
    os.replace(staged_manifest, manifest_path)
    write_forecast_capture_metadata(
        paths.parquet,
        {
            checksum: {
                "retrieved_at_utc": manifest["retrieved_at_utc"],
                "first_seen_capture_at_utc": manifest["retrieved_at_utc"],
                "source_update": manifest["source_update_text"],
            }
        },
    )
    return manifest
