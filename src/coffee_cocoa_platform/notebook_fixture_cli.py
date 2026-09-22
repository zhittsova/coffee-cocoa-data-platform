"""Build and execute the offline synthetic presentation notebook."""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path

import openpyxl

from coffee_cocoa_platform.paths import ProjectPaths
from coffee_cocoa_platform.price_fixture import synthetic_workbook
from coffee_cocoa_platform.prices import publish_workbook
from coffee_cocoa_platform.runtime import run_dbt
from coffee_cocoa_platform.snapshots import open_snapshot, publish_snapshot

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
NOTEBOOK = REPOSITORY_ROOT / "notebooks" / "coffee-cocoa-presentation.ipynb"
START_MONTH = date(2015, 1, 1)
END_MONTH = date(2026, 8, 1)


def _month_sequence() -> list[date]:
    months = []
    month = START_MONTH
    while month <= END_MONTH:
        months.append(month)
        index = month.year * 12 + month.month
        month = date(index // 12, index % 12 + 1, 1)
    return months


def _presentation_workbook() -> bytes:
    """Extend the checked-in synthetic workbook for forecast presentation."""
    book = openpyxl.load_workbook(BytesIO(synthetic_workbook()))
    sheet = book["Monthly Prices"]
    if sheet.max_row > 6:
        sheet.delete_rows(7, sheet.max_row - 6)

    for index, month in enumerate(_month_sequence()):
        cocoa = 10 + 0.03 * index + (month.month % 4) * 0.2
        arabica = 5 + 0.015 * index + (month.month % 3) * 0.1
        robusta = (
            3 + 0.02 * index + (month.month % 5) * 0.08
            if month < date(2026, 7, 1)
            else None
        )
        sheet.append(
            [
                f"{month.year}M{month.month:02}",
                *([None] * 10),
                cocoa,
                arabica,
                robusta,
            ]
        )

    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _build_fixture(root: Path) -> str:
    paths = ProjectPaths.from_root(root)
    if any(
        path.exists()
        for path in (paths.data, paths.warehouse, paths.snapshots, paths.state)
    ):
        raise FileExistsError(
            f"Fixture root already contains project data: {paths.root}; choose an empty root"
        )

    os.environ["COFFEE_COCOA_ENABLE_FORECAST_RESULTS"] = "false"
    os.environ["DAGSTER_DISABLE_TELEMETRY"] = "1"
    from coffee_cocoa_platform.platform_assets import run_fixture_assets

    if not run_fixture_assets(paths.root):
        raise RuntimeError("The offline two-source fixture pipeline failed")

    publish_workbook(
        _presentation_workbook(),
        start=START_MONTH.strftime("%Y-%m"),
        end=END_MONTH.strftime("%Y-%m"),
        paths=paths,
        source_url="synthetic://world-bank-notebook-fixture-v1",
        effective_url="synthetic://world-bank-notebook-fixture-v1",
        retrieved_at=datetime(2026, 9, 23, tzinfo=UTC),
        fixture=True,
    )

    from coffee_cocoa_platform.catalog_metadata import DBT_DIR

    target = paths.state / "dbt-notebook-input"
    build = run_dbt(
        paths.root,
        [
            "build",
            "--full-refresh",
            "--project-dir",
            str(DBT_DIR),
            "--profiles-dir",
            str(DBT_DIR),
            "--target-path",
            str(target),
            "--log-path",
            str(target),
        ],
    )
    if build.returncode:
        raise RuntimeError(
            "The dbt fixture rebuild failed:\n"
            + build.stdout[-4000:]
            + build.stderr[-4000:]
        )

    input_snapshot = publish_snapshot(
        paths.root,
        {"kind": "notebook_fixture_input", "run_id": "s11-notebook-fixture"},
    )
    os.environ["COFFEE_COCOA_HOME"] = str(paths.root)
    os.environ["COFFEE_COCOA_DUCKDB_PATH"] = str(
        paths.warehouse / "coffee_cocoa.duckdb"
    )
    os.environ["COFFEE_COCOA_ENABLE_FORECAST_RESULTS"] = "true"
    os.environ["DAGSTER_DISABLE_TELEMETRY"] = "1"

    from coffee_cocoa_platform.catalog_metadata import ensure_manifest

    ensure_manifest.cache_clear()
    from coffee_cocoa_platform.forecast_assets import run_forecast_assets

    if not run_forecast_assets(paths.root):
        raise RuntimeError("The offline forecast fixture pipeline failed")

    with open_snapshot(paths.root) as reader:
        tables = {row[0] for row in reader.connection.execute("show tables").fetchall()}
        required = {
            "monthly_benchmark_dynamics",
            "monthly_trade_product_metrics",
            "monthly_trade_balances",
            "monthly_partner_concentration",
            "forecast_metrics",
        }
        missing = sorted(required - tables)
        if missing:
            raise RuntimeError(f"The read snapshot is missing marts: {missing}")
        result_snapshot_id = reader.snapshot_id

    return f"input snapshot {input_snapshot['snapshot_id']}; result snapshot {result_snapshot_id}"


def _execute_notebook(root: Path) -> Path:
    jupyter = shutil.which("jupyter")
    if jupyter is None:
        raise RuntimeError(
            "Notebook tools are not installed; run with `uv run --locked --group notebooks`"
        )
    output_name = "coffee-cocoa-presentation-executed.ipynb"
    matplotlib_config = root / ".matplotlib"
    matplotlib_config.mkdir(parents=True, exist_ok=True)
    command = [
        jupyter,
        "nbconvert",
        "--to",
        "notebook",
        "--execute",
        str(NOTEBOOK),
        "--output",
        output_name,
        "--output-dir",
        str(root),
        "--ExecutePreprocessor.timeout=480",
        "--ExecutePreprocessor.kernel_name=python3",
    ]
    environment = dict(
        os.environ,
        COFFEE_COCOA_HOME=str(root),
        DAGSTER_DISABLE_TELEMETRY="1",
        MPLCONFIGDIR=str(matplotlib_config),
    )
    result = subprocess.run(
        command,
        cwd=REPOSITORY_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        timeout=600,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stdout[-4000:] + result.stderr[-4000:])
    return root / output_name


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        required=True,
        help="New or empty directory for synthetic data and executed notebook output",
    )
    args = parser.parse_args()
    root = args.root.expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    if any(root.iterdir()):
        raise SystemExit(f"Choose an empty fixture root: {root}")

    snapshot_summary = _build_fixture(root)
    executed_notebook = _execute_notebook(root)
    print(snapshot_summary)
    print(f"Executed notebook: {executed_notebook}")


if __name__ == "__main__":
    main()
